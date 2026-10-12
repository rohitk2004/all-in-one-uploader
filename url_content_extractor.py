import asyncio
import re
import io
import urllib.request
from html import unescape
from playwright.async_api import async_playwright
import docx

def safe_clean_text(text: str) -> str:
    if not text:
        return ""
    # Remove control characters except newlines/tabs
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)

async def extract_url_content_playwright(url: str) -> dict:
    """Extracts Title and Main Content from a forum/webpage URL using Playwright"""
    url = url.strip()
    if not url.startswith("http://") and not url.startswith("https://"):
        url = "https://" + url

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(
                headless=True,
                args=["--disable-blink-features=AutomationControlled"]
            )
            context = await browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
            )
            page = await context.new_page()

            # Block heavy ad resources
            await page.route(
                "**/*",
                lambda route: route.abort() if any(
                    ad in route.request.url.lower() for ad in [
                        "google-analytics", "doubleclick", "googlesyndication",
                        "quantserve", "rubiconproject", "pubmatic", "criteo", "taboola"
                    ]
                ) else route.continue_()
            )

            await page.goto(url, wait_until="domcontentloaded", timeout=45000)
            await page.wait_for_timeout(2000)

            # Extract Title
            title = ""
            title_el = await page.query_selector(".title-bar h2, .subject, .thread-title, h1.title, h1")
            if title_el:
                title = await title_el.inner_text()
            if not title:
                raw_page_title = await page.title()
                title = raw_page_title.split("|")[0].strip()

            # Extract Content
            content = ""
            # Priority 1: ProBoards / FreeForums first post message
            msg_el = await page.query_selector(".post .message, .thread .message, article.message, div.message")
            if msg_el:
                content = await msg_el.inner_text()
            else:
                # Priority 2: General article or main text
                article_el = await page.query_selector("article, main, .post-content, .entry-content")
                if article_el:
                    content = await article_el.inner_text()
                else:
                    content = await page.inner_text("body")

            await browser.close()
            return {
                "success": True,
                "url": url,
                "title": safe_clean_text(title.strip()),
                "content": safe_clean_text(content.strip())
            }

    except Exception as e:
        # Fallback to urllib parser if Playwright fails
        return extract_url_content_urllib(url)

def extract_url_content_urllib(url: str) -> dict:
    url = url.strip()
    if not url.startswith("http://") and not url.startswith("https://"):
        url = "https://" + url
        
    try:
        req = urllib.request.Request(
            url, 
            headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        )
        with urllib.request.urlopen(req, timeout=25) as response:
            html = response.read().decode('utf-8', errors='ignore')
            
            # Title
            title_match = re.search(r'<title>(.*?)</title>', html, re.DOTALL | re.IGNORECASE)
            raw_title = title_match.group(1).strip() if title_match else ""
            clean_title = unescape(raw_title).split('|')[0].strip()
            
            # Message body
            message_match = re.search(r'<article[^>]*class="[^"]*message[^"]*"[^>]*>(.*?)</article>', html, re.DOTALL | re.IGNORECASE)
            if not message_match:
                message_match = re.search(r'<div[^>]*class="[^"]*message[^"]*"[^>]*>(.*?)</div>', html, re.DOTALL | re.IGNORECASE)
                
            if message_match:
                content_html = message_match.group(1)
                content_html = re.sub(r'<br\s*/?>', '\n', content_html, flags=re.IGNORECASE)
                content_html = re.sub(r'</p>', '\n\n', content_html, flags=re.IGNORECASE)
                content_text = re.sub(r'<[^>]+>', '', content_html)
                clean_content = unescape(content_text).strip()
            else:
                # Strip all tags fallback
                clean_content = unescape(re.sub(r'<[^>]+>', ' ', html)).strip()

            return {
                "success": True,
                "url": url,
                "title": safe_clean_text(clean_title),
                "content": safe_clean_text(clean_content)
            }
    except Exception as err:
        return {
            "success": False,
            "url": url,
            "title": "Extraction Failed",
            "content": "",
            "error": str(err)
        }

def create_docx_file(items: list) -> bytes:
    """
    Creates a Word (.docx) document containing extracted Titles and Contents.
    items: list of dicts with 'url', 'title', 'content'
    """
    doc = docx.Document()
    
    # Set document title
    doc.add_heading("Extracted Forum & Web Contents", level=0)
    doc.add_paragraph(f"Total Extracted URLs: {len(items)}\n")
    
    for idx, item in enumerate(items, start=1):
        if not item.get("success"):
            continue
            
        title = item.get("title", f"Article {idx}")
        content = item.get("content", "")
        url = item.get("url", "")

        # Article Heading
        doc.add_heading(f"{idx}. {title}", level=1)
        if url:
            p_url = doc.add_paragraph()
            p_url.add_run(f"Source URL: {url}").italic = True
        
        # Article Content Paragraphs
        for paragraph in content.split("\n\n"):
            p_text = paragraph.strip()
            if p_text:
                doc.add_paragraph(p_text)
                
        doc.add_paragraph("\n" + "_"*50 + "\n")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()

def create_single_docx(title: str, content: str, url: str = "") -> bytes:
    """Creates a .docx file for a single extracted item"""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    if url:
        p = doc.add_paragraph()
        p.add_run(f"Source: {url}").italic = True
    
    for para in content.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
            
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream.getvalue()
