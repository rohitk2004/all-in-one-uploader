import docx
import re
import io

def parse_docx_articles(docx_source) -> list:
    """
    Parses a .docx file containing multiple articles formatted with headings or title lines.
    Returns a list of dicts: [{'title': ..., 'content': ...}]
    """
    if isinstance(docx_source, bytes):
        doc = docx.Document(io.BytesIO(docx_source))
    elif isinstance(docx_source, str):
        doc = docx.Document(docx_source)
    elif isinstance(docx_source, docx.document.Document):
        doc = docx_source
    else:
        # File-like object (e.g. Streamlit UploadedFile)
        file_bytes = docx_source.read() if hasattr(docx_source, "read") else docx_source
        doc = docx.Document(io.BytesIO(file_bytes))

    articles = []
    current_title = ""
    current_paras = []

    def save_current_article():
        nonlocal current_title, current_paras
        if current_title and current_paras:
            full_content = "\n\n".join(current_paras).strip()
            # Filter out Source URL lines if present
            lines = full_content.splitlines()
            clean_lines = [l for l in lines if not l.strip().lower().startswith("source url:")]
            clean_content = "\n".join(clean_lines).strip()
            
            if clean_content:
                articles.append({
                    "title": current_title.strip(),
                    "content": clean_content
                })
        current_title = ""
        current_paras = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        # Check if paragraph is a Heading style or starts with a numbered heading like "1. Title"
        is_heading = p.style.name.startswith("Heading") or bool(re.match(r'^\d+\.\s+[\w\W]+', text))
        
        # Check horizontal divider lines
        if text.startswith("____"):
            save_current_article()
            continue

        if is_heading:
            save_current_article()
            clean_t = re.sub(r'^\d+\.\s*', '', text).strip()
            current_title = clean_t
        else:
            if not current_title:
                current_title = text
            else:
                current_paras.append(text)

    # Save final article
    save_current_article()
    return articles
